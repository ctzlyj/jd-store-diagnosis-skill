# -*- coding: utf-8 -*-
"""Structural gate for a shop diagnosis deliverable (markdown + Excel).

Usage:
  python -X utf8 scripts/validate_report.py \
    --md "报告/某店-十月经营诊断报告.md" \
    --xlsx "报告/某店-十月经营数据.xlsx" \
    --shop-id 204560127 --shop-name 东来臻选经营部

Exit code 0 only when every check passes. Do not replace this with eyeballing.
"""
import argparse
import os
import re
import sys


def read_text(path):
    raw = open(path, "rb").read()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16")
    return raw.decode("utf-8", errors="replace")


def check(ok, name, detail=""):
    print(("PASS" if ok else "FAIL"), "|", name, ("| " + detail if detail and not ok else ""))
    return ok


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--shop-id", required=True)
    ap.add_argument("--shop-name", required=True)
    ap.add_argument("--min-chars", type=int, default=1500)
    args = ap.parse_args()

    fails = []

    if not os.path.exists(args.md):
        print("FAIL | md 文件存在 |", args.md)
        sys.exit(1)
    md = read_text(args.md)
    fails.append(not check(len(md) >= args.min_chars, "md 正文长度",
                           f"{len(md)} < {args.min_chars}"))

    title = md.splitlines()[0] if md.splitlines() else ""
    fails.append(not check(title.startswith("# ") and args.shop_name in title,
                           "标题带店铺名", title[:80]))

    id_hits = re.findall(r"店铺ID[：:]\s*(\d{6,})", md)
    fails.append(not check(len(id_hits) >= 2, "店铺ID 至少出现两次（开头+数据说明）",
                           f"出现 {len(id_hits)} 次"))
    fails.append(not check(all(x == args.shop_id for x in id_hits), "店铺ID 与参数一致",
                           f"md={id_hits} arg={args.shop_id}"))

    for section in ("一句话结论", "行动清单", "核心经营数据", "数据说明"):
        fails.append(not check(section in md, f"章节·{section}"))

    tiers = ("立即" in md and "本周" in md and "本月" in md) or \
            (re.search(r"^###\s*A[.、]", md, re.M) and re.search(r"^###\s*B[.、]", md, re.M)
             and re.search(r"^###\s*C[.、]", md, re.M))
    fails.append(not check(bool(tiers), "行动清单含 A/B/C 或 立即/本周/本月"))

    for key in ("数据来源", "期间", "口径"):
        fails.append(not check(key in md, f"数据说明·{key}"))

    fails.append(not check("~" not in md, "不含 ~ 区间（JoySpace 会渲染成删除线）"))

    placeholders = [p for p in ("TODO", "待补", "XXX", "{店铺名}", "待填写") if p in md]
    fails.append(not check(not placeholders, "无占位符", "、".join(placeholders)))

    if not os.path.exists(args.xlsx):
        print("FAIL | xlsx 文件存在 |", args.xlsx)
        sys.exit(1)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    except Exception as exc:
        print("FAIL | xlsx 可打开 |", str(exc)[:120])
        sys.exit(1)

    fails.append(not check(len(wb.sheetnames) >= 3, "Excel 至少 3 张表",
                           "、".join(wb.sheetnames)))
    fails.append(not check(any("数据说明" in s for s in wb.sheetnames), "Excel 含数据说明表"))

    first = wb[wb.sheetnames[0]]
    id_row = None
    for row in first.iter_rows(min_row=1, max_row=min(first.max_row, 20), max_col=4):
        if row[0].value == "店铺ID":
            id_row = row
            break
    ok_id = bool(id_row) and str(id_row[1].value).strip() == args.shop_id
    fails.append(not check(ok_id, "Excel 核心表含店铺ID行且值一致",
                           f"got={id_row[1].value if id_row else None}"))

    labels = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    labels.add(value)
    for key in ("成交金额", "访客", "客单价"):
        fails.append(not check(any(key in x for x in labels), f"Excel 含指标·{key}"))

    fails.append(not check(args.shop_name in md and args.shop_name in labels,
                           "店铺名在 md 与 Excel 中一致"))

    print("=" * 46)
    if any(fails):
        print(f"共 {sum(fails)} 项 FAIL：结构不达标，先补齐再发布")
        sys.exit(1)
    print("结论：结构自检全部通过，可发布")


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    main()
